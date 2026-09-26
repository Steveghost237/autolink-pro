# -*- coding: utf-8 -*-
"""
AutoLink Pro — Présentation simple du fonctionnement (non technique)
Génère : AutoLink_Presentation_Simple.docx
"""
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import datetime

TODAY = datetime.date.today().strftime('%d %B %Y')
OUT = r'c:\Users\Albert WIB\CascadeProjects\autolink'

C_DARK = RGBColor(0x0F, 0x17, 0x2A)
C_TEAL = RGBColor(0x0D, 0x94, 0x88)
C_GREY = RGBColor(0x64, 0x74, 0x8B)
C_WHT  = RGBColor(0xFF, 0xFF, 0xFF)
C_GRN  = RGBColor(0x10, 0xB9, 0x81)
C_AMB  = RGBColor(0xB4, 0x53, 0x09)


def shd(cell, hex6):
    tc = cell._tc
    pr = tc.get_or_add_tcPr()
    e = OxmlElement('w:shd')
    e.set(qn('w:val'), 'clear'); e.set(qn('w:color'), 'auto'); e.set(qn('w:fill'), hex6)
    pr.append(e)


def new_doc():
    doc = Document()
    s = doc.sections[0]
    s.page_width = Cm(21); s.page_height = Cm(29.7)
    s.left_margin = Cm(2.5); s.right_margin = Cm(2.5)
    s.top_margin = Cm(2.5); s.bottom_margin = Cm(2.5)
    return doc


def h1(doc, text):
    p = doc.add_heading('', level=1)
    r = p.add_run(text)
    r.font.color.rgb = C_DARK
    r.font.size = Pt(18)
    p.space_after = Pt(6)
    return p


def h2(doc, text):
    p = doc.add_heading('', level=2)
    r = p.add_run(text)
    r.font.color.rgb = C_TEAL
    r.font.size = Pt(14)
    return p


def para(doc, text, bold=False, color=None, size=11, space_after=6):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = bold
    r.font.size = Pt(size)
    if color:
        r.font.color.rgb = color
    p.paragraph_format.space_after = Pt(space_after)
    return p


def bullet(doc, text, bold_prefix=None):
    p = doc.add_paragraph(style='List Bullet')
    if bold_prefix:
        r = p.add_run(bold_prefix)
        r.bold = True
        r.font.size = Pt(11)
    r = p.add_run(text)
    r.font.size = Pt(11)
    return p


def step_table(doc, rows):
    t = doc.add_table(rows=len(rows) + 1, cols=3)
    t.style = 'Table Grid'
    hdr = t.rows[0].cells
    for i, txt in enumerate(['Étape', 'Qui agit ?', 'Ce qui se passe']):
        shd(hdr[i], '0D9488')
        p = hdr[i].paragraphs[0]
        r = p.add_run(txt)
        r.bold = True
        r.font.color.rgb = C_WHT
        r.font.size = Pt(10)
    for i, (n, qui, quoi) in enumerate(rows):
        cells = t.rows[i + 1].cells
        cells[0].paragraphs[0].add_run(str(n)).bold = True
        cells[1].paragraphs[0].add_run(qui).font.size = Pt(10)
        cells[2].paragraphs[0].add_run(quoi).font.size = Pt(10)
        for c in cells:
            for p in c.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(10)
    return t


def simple_table(doc, header, rows, widths=None):
    t = doc.add_table(rows=len(rows) + 1, cols=len(header))
    t.style = 'Table Grid'
    hdr = t.rows[0].cells
    for i, txt in enumerate(header):
        shd(hdr[i], '0F172A')
        r = hdr[i].paragraphs[0].add_run(txt)
        r.bold = True
        r.font.color.rgb = C_WHT
        r.font.size = Pt(10)
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            p = t.rows[i + 1].cells[j].paragraphs[0]
            r = p.add_run(str(val))
            r.font.size = Pt(10)
    return t


def main():
    doc = new_doc()

    # ── Page de garde ──────────────────────────────────────────────
    for _ in range(6):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run('AutoLink Pro')
    r.font.size = Pt(40)
    r.bold = True
    r.font.color.rgb = C_TEAL

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run('La location de voitures au Cameroun,\nsimplement, depuis votre téléphone')
    r.font.size = Pt(16)
    r.font.color.rgb = C_DARK

    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run('Document de présentation — fonctionnement expliqué simplement')
    r.font.size = Pt(11)
    r.font.color.rgb = C_GREY
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(f'Douala — {TODAY}')
    r.font.size = Pt(10)
    r.font.color.rgb = C_GREY

    doc.add_page_break()

    # ── 1. C'est quoi ? ────────────────────────────────────────────
    h1(doc, '1. C’est quoi AutoLink Pro ?')
    para(doc,
         "AutoLink Pro est une application de location de voitures pensée pour le Cameroun. "
         "Imaginez une grande agence de location — sauf qu'elle tient entièrement dans un "
         "téléphone et un site web.")
    bullet(doc, "recherchent une voiture, paient et reçoivent leur réservation.", bold_prefix="Les clients ")
    bullet(doc, "mettent leur voiture en location et gagnent de l'argent.", bold_prefix="Les propriétaires ")
    bullet(doc, "conduisent les clients qui le souhaitent.", bold_prefix="Les chauffeurs ")
    bullet(doc, "surveille tout depuis son écran et arbitre en cas de problème.", bold_prefix="L'administrateur ")
    para(doc,
         "Tout le monde est connecté au même système : une réservation faite sur le téléphone "
         "d'un client apparaît immédiatement sur l'écran de l'administrateur — et inversement.",
         color=C_TEAL)

    # ── 2. Les 5 profils ───────────────────────────────────────────
    h1(doc, '2. Qui utilise l’application ?')
    simple_table(doc,
        ['Profil', 'Ce qu’il fait'],
        [
            ['CLIENT', "Cherche une voiture, choisit ses dates, paie, utilise le véhicule, le rend."],
            ['PROPRIÉTAIRE', "Inscrit sa voiture, fixe son prix dans la fourchette autorisée, encaisse sa part (50 %)."],
            ['CHAUFFEUR', "Reçoit ses courses assignées, conduit le client, marque la fin de course. Son compte est créé par l'admin."],
            ['CONTRÔLEUR', "Inspecte les voitures au départ et au retour (état, kilométrage, photos)."],
            ['ADMINISTRATEUR', "Supervise tout : valide les voitures, arbitre les litiges, gère les comptes du personnel."],
        ])
    para(doc,
         "Important : sur l'application mobile, seuls les clients et les propriétaires peuvent "
         "créer un compte. Les comptes chauffeurs et contrôleurs sont créés par l'administrateur — "
         "le personnel reste sous son contrôle.", size=10, color=C_GREY)

    # ── 3. Comment se passe une location ? ─────────────────────────
    h1(doc, '3. Comment se passe une location, du début à la fin ?')
    para(doc, "Voici le parcours complet, étape par étape :")
    step_table(doc, [
        (1, 'Client', "Ouvre l'app, parcourt le catalogue, choisit une voiture, ses dates, "
                      "et s'il veut un chauffeur ou non."),
        (2, 'Client', "Paie : solde AutoLink, MTN MoMo, Orange Money, SenBid, PayBid ou carte bancaire. "
                      "Le prix affiché inclut la location + une caution restituable."),
        (3, 'Le système', "La réservation est confirmée AUTOMATIQUEMENT dès que le paiement passe. "
                          "Aucun clic de l'admin n'est nécessaire."),
        (4, 'Le système', "Le propriétaire est notifié. Si un chauffeur AutoLink est demandé, "
                          "il est assigné automatiquement."),
        (5, 'Administrateur', "Voit la réservation apparaître sur son écran en quelques secondes — "
                              "il supervise, il n'a plus à valider à la main."),
        (6, 'Client', "Prend la voiture et l'utilise (avec chauffeur si choisi)."),
        (7, 'Chauffeur ou proprio', "Marque la location « terminée » au retour du véhicule "
                                    "(ou elle se termine toute seule à la date prévue)."),
        (8, 'Le système', "La voiture redevient IMMÉDIATEMENT disponible dans le catalogue — "
                          "un autre client peut la louer tout de suite."),
        (9, 'Le système', "L'argent est réparti : 50 % pour AutoLink, 50 % pour le propriétaire. "
                          "La caution du client est restituée."),
    ])

    # ── 4. L'argent ────────────────────────────────────────────────
    h1(doc, '4. Comment circule l’argent ?')
    para(doc, "Prenons un exemple : une location à 100 000 FCFA.", bold=True)
    bullet(doc, "reçoivent 50 000 F (commission).", bold_prefix="AutoLink ")
    bullet(doc, "reçoit 50 000 F — mais bloqués en CAUTION pendant la location. "
                "Ils ne lui sont versés qu'au retour de la voiture, en bon état.",
           bold_prefix="Le propriétaire ")
    bullet(doc, "paie en plus une caution de garantie selon la catégorie "
                "(50 000 à 800 000 F), restituée au retour.", bold_prefix="Le client ")
    para(doc,
         "Si le client signale un problème (panne, dégât) : il ouvre un LITIGE. "
         "L'argent reste gelé, et l'administrateur arbitre — soit il rembourse le client, "
         "soit il verse au propriétaire.", color=C_AMB)

    # ── 5. Le catalogue ────────────────────────────────────────────
    h1(doc, '5. Les voitures : 5 catégories, attribuées automatiquement')
    para(doc,
         "Le propriétaire ne choisit PAS la catégorie de sa voiture. Le système la calcule "
         "tout seul à partir de faits vérifiables : la marque, le modèle, l'année, le kilométrage, "
         "l'état et la valeur de la voiture sur le marché.")
    simple_table(doc,
        ['Catégorie', 'Exemples', 'Prix indicatif / jour', 'Caution'],
        [
            ['Économique', 'Toyota Yaris, Kia Picanto, Renault Logan', '15 000 – 30 000 F', '50 000 F'],
            ['Intermédiaire', 'Toyota Corolla, Hyundai Tucson, Dacia Duster', '30 000 – 60 000 F', '100 000 F'],
            ['Premium', 'Toyota Prado, Mercedes Classe E, BMW X5', '60 000 – 110 000 F', '200 000 F'],
            ['Luxe', 'Mercedes GLE, Range Rover, Classe S', '110 000 – 180 000 F', '400 000 F'],
            ['Super Luxe', 'Bentley, Porsche, Maserati', '180 000 F et +', '800 000 F'],
        ])
    bullet(doc, "Le prix de référence s'adapte à la ville (Douala = référence, Yaoundé −5 %, autres villes −10 à −15 %).")
    bullet(doc, "Location de 7 jours et plus : −5 % automatique. 30 jours et plus : −10 %.")
    bullet(doc, "Chaque catégorie inclut un forfait kilométrique (200 à 350 km/jour) et un prix par km supplémentaire.")
    bullet(doc, "Le propriétaire peut ajuster son prix, mais uniquement dans une fourchette de ±20 % "
                "autour du prix indicatif — impossible de surcharger ou de casser le marché.")

    # ── 6. Le portefeuille ─────────────────────────────────────────
    h1(doc, '6. Le portefeuille AutoLink')
    para(doc,
         "Chaque client a un solde dans l'application, qu'il recharge avec : "
         "MTN MoMo, Orange Money, SenBid, PayBid, PayPal ou carte bancaire (Stripe). "
         "Il paie ensuite ses locations depuis ce solde — comme du crédit téléphonique.")

    # ── 7. Où ça se passe ──────────────────────────────────────────
    h1(doc, '7. Les trois portes d’entrée')
    simple_table(doc,
        ['Porte', 'Pour qui', 'Adresse / support'],
        [
            ['Application mobile (APK)', 'Clients et propriétaires', 'Téléchargement Android depuis le site'],
            ['Site web', 'Tout le monde + administration', 'autolink-pro.worldwide-international.business'],
            ['Le « cerveau » (serveur)', 'Invisible — connecte tout', 'api-autolink-pro.worldwide-international.business'],
        ])
    para(doc,
         "Les deux portes (mobile et web) parlent au même cerveau : toutes les réservations, "
         "comptes et paiements sont au même endroit. C'est ce qui rend la synchronisation "
         "instantanée possible.")

    # ── 8. Résumé en une phrase ────────────────────────────────────
    h1(doc, '8. En une phrase')
    para(doc,
         "AutoLink Pro connecte trois mondes : ceux qui ont besoin d'une voiture, "
         "ceux qui en ont une à louer, et ceux qui savent conduire — le tout orchestré "
         "automatiquement, sécurisé par des cautions, et supervisé par un administrateur "
         "qui voit tout en temps réel.",
         bold=True, color=C_TEAL, size=12)

    doc.save(OUT + r'\AutoLink_Presentation_Simple.docx')
    print('OK ->', OUT + r'\AutoLink_Presentation_Simple.docx')


if __name__ == '__main__':
    main()
